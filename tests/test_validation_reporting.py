"""A baseline must not hide the defects still present in a deliverable."""

import contextlib
import copy
import io
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
OFFICE = ROOT / "skills/pptx/scripts/office"
sys.path.insert(0, str(OFFICE))
from validators.pptx import PPTXSchemaValidator
import pack

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


class ValidationReporting(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.tree = self.root / "parts"
        self.source = self.root / "source.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Test"
        slide.notes_slide.notes_text_frame.text = "Notes belonging to this slide"
        prs.save(self.source)
        with zipfile.ZipFile(self.source) as archive:
            archive.extractall(self.tree)

    def mutate(self, relative, callback):
        path = self.tree / relative
        tree = etree.parse(str(path))
        callback(tree.getroot())
        tree.write(str(path), encoding="utf-8", xml_declaration=True)

    def snapshot(self, name):
        path = self.root / name
        with zipfile.ZipFile(path, "w") as archive:
            for part in self.tree.rglob("*"):
                if part.is_file():
                    archive.write(part, part.relative_to(self.tree))
        return path

    def validate(self, baseline=None):
        validator = PPTXSchemaValidator(self.tree, baseline)
        with contextlib.redirect_stdout(io.StringIO()):
            ok = validator.validate()
        return ok, validator, validator.validation_summary(ok)

    def negative_extent(self):
        self.mutate("ppt/slides/slide1.xml", lambda root: root.find(f".//{{{A}}}ext").set("cx", "-100"))

    def test_inherited_negative_extent_not_reported_as_clean(self):
        self.negative_extent()
        baseline = self.snapshot("invalid-baseline.pptx")
        ok, validator, summary = self.validate(baseline)
        self.assertTrue(ok)  # Accept baseline-preserving template edits.
        result = validator.xsd_results["ppt/slides/slide1.xml"]
        self.assertTrue(result["inherited"])
        self.assertFalse(result["new"])
        self.assertIn("NO_NEW_ERRORS", summary)
        self.assertNotIn("All validations PASSED", summary)
        self.assertFalse(self.validate()[0])  # Same artifact fails full validation.

    def test_repeated_error_instances_expose_baseline_message_set_limit(self):
        # Two separate objects, only one invalid in the baseline. A later second
        # instance has the same schema diagnostic, so set-based acceptance alone
        # cannot establish that the number or position of defects is preserved.
        def add_second_shape(root):
            original = root.find(f".//{{{P}}}sp")
            second = copy.deepcopy(original)
            second.find(f".//{{{P}}}cNvPr").set("id", "3")
            second.find(f".//{{{P}}}cNvPr").set("name", "Second textbox")
            original.getparent().append(second)
        self.mutate("ppt/slides/slide1.xml", add_second_shape)
        self.negative_extent()
        baseline = self.snapshot("one-invalid-object.pptx")
        _, original_validator, _ = self.validate(baseline)
        original_messages = original_validator.xsd_results["ppt/slides/slide1.xml"]["current"]
        self.mutate(
            "ppt/slides/slide1.xml",
            lambda root: root.findall(f".//{{{P}}}sp")[1].find(f".//{{{A}}}ext").set("cx", "-100"),
        )
        root = etree.parse(str(self.tree / "ppt/slides/slide1.xml"))
        self.assertEqual(len(root.xpath('//*[local-name()="ext"][@cx="-100"]')), 2)
        ok, validator, summary = self.validate(baseline)
        result = validator.xsd_results["ppt/slides/slide1.xml"]
        self.assertTrue(ok)
        self.assertEqual(result["current"], original_messages)
        self.assertFalse(result["new"])
        self.assertIn("NO_NEW_ERRORS", summary)
        self.assertIn("no new distinct XSD diagnostic messages per part", summary)
        self.assertIn("repeated-error counts and locations are not compared", summary)
        self.assertFalse(self.validate()[0])

    def test_new_error_fails_even_with_inherited_error(self):
        self.negative_extent()
        baseline = self.snapshot("invalid-baseline.pptx")
        self.mutate("ppt/presentation.xml", lambda root: root.insert(0, etree.Element(f"{{{P}}}invalid")))
        ok, validator, summary = self.validate(baseline)
        self.assertFalse(ok)
        self.assertTrue(validator.xsd_results["ppt/slides/slide1.xml"]["inherited"])
        self.assertTrue(validator.xsd_results["ppt/presentation.xml"]["new"])
        self.assertIn("FAILED", summary)

    def test_fixed_error_is_not_still_counted_as_inherited(self):
        self.negative_extent()
        baseline = self.snapshot("invalid-baseline.pptx")
        self.mutate("ppt/slides/slide1.xml", lambda root: root.find(f".//{{{A}}}ext").set("cx", "100"))
        ok, validator, summary = self.validate(baseline)
        self.assertTrue(ok)
        self.assertFalse(validator.xsd_results["ppt/slides/slide1.xml"]["current"])
        self.assertNotIn("NO_NEW_ERRORS", summary)

    def test_skipped_parts_are_reported_without_full_pass_claim(self):
        path = self.tree / "customXml/item1.xml"
        path.parent.mkdir()
        path.write_text('<metadata xmlns="urn:fixture"/>')
        validator = PPTXSchemaValidator(self.tree)
        with contextlib.redirect_stdout(io.StringIO()):
            ok = validator.validate_against_xsd()
        summary = validator.validation_summary(ok)
        self.assertTrue(ok)
        self.assertTrue(validator.xsd_results["customXml/item1.xml"]["skipped"])
        self.assertIn("skipped (no schema): customXml/item1.xml", summary)
        self.assertIn("WITH_LIMITATIONS", summary)

    def test_ignored_diagnostics_remain_visible(self):
        # The bundled core-properties XSD has an unresolved dcterms reference.
        # This real, pre-existing validator limitation must be disclosed.
        ok, validator, summary = self.validate()
        self.assertTrue(ok)  # Existing ignore policy is retained, not concealed.
        result = validator.xsd_results["docProps/core.xml"]
        self.assertTrue(result["current"])
        self.assertTrue(result["ignored"])
        self.assertTrue(any("purl.org/dc/terms" in e for e in result["ignored"]))
        self.assertFalse(result["new"])
        self.assertIn("WITH_LIMITATIONS", summary)
        self.assertIn("ignored: docProps/core.xml", summary)

    def test_notes_master_order_violation_is_detected(self):
        def misorder(root):
            notes = root.find(f"{{{P}}}notesMasterIdLst")
            slides = root.find(f"{{{P}}}sldIdLst")
            if notes is None:
                notes = etree.Element(f"{{{P}}}notesMasterIdLst")
                etree.SubElement(notes, f"{{{P}}}notesMasterId", attrib={
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id": "rId1"
                })
            else:
                root.remove(notes)
            root.insert(root.index(slides) + 1, notes)
        self.mutate("ppt/presentation.xml", misorder)
        validator = PPTXSchemaValidator(self.tree)
        with contextlib.redirect_stdout(io.StringIO()):
            ok = validator.validate_against_xsd()
        self.assertFalse(ok)
        self.assertTrue(any("notesMasterIdLst" in err for err in validator.xsd_results["ppt/presentation.xml"]["new"]))

    def test_cli_and_pack_disclose_inherited_defects(self):
        self.negative_extent()
        baseline = self.snapshot("invalid-baseline.pptx")
        run = subprocess.run([sys.executable, str(OFFICE / "validate.py"), str(baseline), "--original", str(baseline)], capture_output=True, text=True)
        self.assertEqual(run.returncode, 0, run.stdout + run.stderr)
        self.assertIn("NO_NEW_ERRORS", run.stdout)
        self.assertIn("inherited: ppt/slides/slide1.xml", run.stdout)
        output = self.root / "packed.pptx"
        stream = io.StringIO()
        with contextlib.redirect_stdout(stream):
            _, message = pack.pack(str(self.tree), str(output), str(baseline))
        self.assertTrue(output.exists(), message)
        self.assertIn("NO_NEW_ERRORS", stream.getvalue())
        self.assertNotIn("All validations PASSED", stream.getvalue())


if __name__ == "__main__":
    unittest.main()
