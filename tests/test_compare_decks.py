"""Edit-preservation checks use real OPC packages and observable mutations."""
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills/pptx/scripts"))
from compare_decks import compare

N = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
     "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


class DeckComparisonTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.before = self.root / "before.pptx"
        self.after = self.root / "after.pptx"
        prs = Presentation()
        for i in range(3):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
            shape.text = f"Research step {i + 1}"
            shape.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(10, 20, 30)
            slide.notes_slide.notes_text_frame.text = f"Speaker notes for step {i + 1}"
        prs.save(self.before)
        self.after.write_bytes(self.before.read_bytes())

    def mutate(self, path, action, source=None, destination=None):
        source = source or self.after
        destination = destination or self.after
        with zipfile.ZipFile(source) as archive:
            parts = {n: archive.read(n) for n in archive.namelist()}
        root = etree.fromstring(parts[path])
        action(root)
        parts[path] = etree.tostring(root)
        with zipfile.ZipFile(destination, "w") as archive:
            for name, data in parts.items():
                archive.writestr(name, data)

    @staticmethod
    def reorder(root):
        order = root.find("p:sldIdLst", N)
        first = order[0]
        order.remove(first)
        order.append(first)

    def test_unchanged_deck_passes_all_modes(self):
        for mode in ("palette", "layout", "reorder"):
            with self.subTest(mode=mode):
                self.assertTrue(compare(self.before, self.after, mode)["passed"])

    def test_palette_change_preserves_text_notes_and_geometry(self):
        self.mutate("ppt/slides/slide1.xml", lambda root: root.xpath('//a:srgbClr', namespaces=N)[0].set("val", "275A92"))
        report = compare(self.before, self.after, "palette", [1])
        self.assertTrue(report["passed"], report)
        self.assertEqual(report["changed_parts"], ["ppt/slides/slide1.xml"])

    def test_palette_rejects_text_change(self):
        def edit(root):
            root.xpath('//a:t', namespaces=N)[0].text = "Changed claim"
        self.mutate("ppt/slides/slide1.xml", edit)
        self.assertFalse(compare(self.before, self.after, "palette")["passed"])

    def test_palette_rejects_font_size_change(self):
        self.mutate("ppt/slides/slide1.xml", lambda root: root.xpath('//a:rPr', namespaces=N)[0].set("sz", "1000"))
        self.assertFalse(compare(self.before, self.after, "palette")["passed"])

    def test_layout_allows_geometry_but_palette_rejects_it(self):
        self.mutate("ppt/slides/slide1.xml", lambda root: root.xpath('//a:xfrm/a:off', namespaces=N)[0].set("x", "1500000"))
        self.assertTrue(compare(self.before, self.after, "layout", [1])["passed"])
        self.assertFalse(compare(self.before, self.after, "palette", [1])["passed"])

    def test_layout_rejects_text_mutation(self):
        def edit(root):
            root.xpath('//a:t', namespaces=N)[0].text = "Changed claim"
        self.mutate("ppt/slides/slide1.xml", edit)
        self.assertFalse(compare(self.before, self.after, "layout", [1])["passed"])

    def test_unselected_slide_change_is_rejected(self):
        self.mutate("ppt/slides/slide2.xml", lambda root: root.xpath('//a:srgbClr', namespaces=N)[0].set("val", "275A92"))
        report = compare(self.before, self.after, "palette", [1])
        self.assertFalse(report["passed"])
        self.assertTrue(any("Unselected slide 2" in f for f in report["failures"]))

    def test_selected_positions_follow_display_order(self):
        self.mutate("ppt/presentation.xml", self.reorder, source=self.before, destination=self.before)
        self.after.write_bytes(self.before.read_bytes())
        self.mutate("ppt/slides/slide2.xml", lambda root: root.xpath('//a:srgbClr', namespaces=N)[0].set("val", "275A92"))
        report = compare(self.before, self.after, "palette", [1])
        self.assertTrue(report["passed"], report)
        self.assertEqual(report["before_order"][0], "ppt/slides/slide2.xml")

    def test_reorder_keeps_notes_attached_to_slide_identity(self):
        self.mutate("ppt/presentation.xml", self.reorder)
        report = compare(self.before, self.after, "reorder")
        self.assertTrue(report["passed"], report)
        self.assertEqual(report["after_order"], ["ppt/slides/slide2.xml", "ppt/slides/slide3.xml", "ppt/slides/slide1.xml"])
        self.assertFalse(compare(self.before, self.after, "palette")["passed"])

    def test_swapped_notes_relationship_fails(self):
        self.mutate("ppt/presentation.xml", self.reorder)
        def swap(root):
            for rel in root:
                if rel.get("Type", "").endswith("/notesSlide"):
                    rel.set("Target", "../notesSlides/notesSlide2.xml")
        self.mutate("ppt/slides/_rels/slide1.xml.rels", swap)
        report = compare(self.before, self.after, "reorder")
        self.assertFalse(report["passed"])
        self.assertTrue(any("notes_part changed" in f for f in report["failures"]))

    def test_notes_content_change_is_detected(self):
        def edit(root):
            body = root.xpath('//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]', namespaces=N)[0]
            body.xpath('.//a:t', namespaces=N)[0].text = "Wrong slide notes"
        self.mutate("ppt/notesSlides/notesSlide1.xml", edit)
        self.assertFalse(compare(self.before, self.after, "layout", [1])["passed"])

    def test_shared_theme_change_affects_unselected_slides(self):
        self.mutate("ppt/theme/theme1.xml", lambda root: root.xpath('//a:srgbClr', namespaces=N)[0].set("val", "275A92"))
        report = compare(self.before, self.after, "palette", [1])
        self.assertFalse(report["passed"])
        self.assertTrue(any("Unselected slide 2" in f for f in report["failures"]))
        self.assertTrue(compare(self.before, self.after, "palette")["passed"])

    def test_invalid_scope_and_missing_internal_part_fail(self):
        with self.assertRaises(ValueError):
            compare(self.before, self.after, "layout", [4])
        def edit(root):
            for rel in root:
                if rel.get("Type", "").endswith("/notesSlide"):
                    rel.set("Target", "../../../../outside.xml")
        self.mutate("ppt/slides/_rels/slide1.xml.rels", edit)
        with self.assertRaises(ValueError):
            compare(self.before, self.after, "layout")

    def test_layout_rejects_changed_embedded_image(self):
        from PIL import Image
        image = self.root / "sample.png"
        Image.new("RGB", (20, 10), "red").save(image)
        prs = Presentation(self.before)
        prs.slides[0].shapes.add_picture(str(image), Inches(1), Inches(3))
        prs.save(self.before)
        with zipfile.ZipFile(self.before) as archive:
            parts = {n: archive.read(n) for n in archive.namelist()}
        media = next(n for n in parts if n.startswith("ppt/media/"))
        Image.new("RGB", (20, 10), "blue").save(image)
        parts[media] = image.read_bytes()
        with zipfile.ZipFile(self.after, "w") as archive:
            for name, data in parts.items():
                archive.writestr(name, data)
        report = compare(self.before, self.after, "layout", [1])
        self.assertFalse(report["passed"])
        self.assertTrue(any("Embedded content changed" in f for f in report["failures"]))

    def test_global_default_color_rejects_partial_scope(self):
        def edit(root):
            root.xpath('./p:defaultTextStyle//a:schemeClr', namespaces=N)[0].set('val', 'accent2')
        self.mutate('ppt/presentation.xml', edit)
        self.assertFalse(compare(self.before, self.after, 'palette', [1])['passed'])
        self.assertTrue(compare(self.before, self.after, 'palette')['passed'])

    def test_global_table_style_rejects_partial_scope(self):
        # Install a global custom style and make tables on two slides use it.
        style_id = '{12345678-1234-1234-1234-123456789012}'
        prs = Presentation(self.before)
        for slide in list(prs.slides)[:2]:
            table = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(2), Inches(1))
            table.table.cell(0, 0).text = 'Shared table style'
            table._element.xpath('.//a:tableStyleId')[0].text = style_id
        prs.save(self.before)
        def install(root):
            style = etree.SubElement(root, '{' + N['a'] + '}tblStyle', styleId=style_id, styleName='Shared')
            cell = etree.SubElement(etree.SubElement(style, '{' + N['a'] + '}wholeTbl'), '{' + N['a'] + '}tcStyle')
            fill = etree.SubElement(etree.SubElement(cell, '{' + N['a'] + '}fill'), '{' + N['a'] + '}solidFill')
            etree.SubElement(fill, '{' + N['a'] + '}srgbClr', val='123456')
        self.mutate('ppt/tableStyles.xml', install, source=self.before, destination=self.before)
        self.after.write_bytes(self.before.read_bytes())
        self.mutate('ppt/tableStyles.xml', lambda r: r.xpath('//a:srgbClr', namespaces=N)[0].set('val', 'ABCDEF'))
        self.assertFalse(compare(self.before, self.after, 'palette', [1])['passed'])
        self.assertTrue(compare(self.before, self.after, 'palette')['passed'])

    def add_two_images(self):
        from PIL import Image
        prs = Presentation(self.before)
        for i, color in enumerate(('red', 'blue')):
            path = self.root / f'{color}.png'
            Image.new('RGB', (20, 10), color).save(path)
            prs.slides[0].shapes.add_picture(str(path), Inches(1 + i), Inches(3))
        prs.save(self.before)
        self.after.write_bytes(self.before.read_bytes())

    def test_layout_rejects_image_target_change(self):
        self.add_two_images()
        def edit(root):
            images = [r for r in root if r.get('Type', '').endswith('/image')]
            images[0].set('Target', images[1].get('Target'))
        self.mutate('ppt/slides/_rels/slide1.xml.rels', edit)
        self.assertFalse(compare(self.before, self.after, 'layout', [1])['passed'])

    def test_layout_rejects_image_reference_swap_but_allows_moving(self):
        self.add_two_images()
        def move(root):
            root.xpath('//p:pic//a:off', namespaces=N)[0].set('x', '999999')
        self.mutate('ppt/slides/slide1.xml', move)
        self.assertTrue(compare(self.before, self.after, 'layout', [1])['passed'])
        def swap(root):
            nodes = root.xpath('//a:blip', namespaces=N)
            key = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'
            x, y = [n.get(key) for n in nodes]
            nodes[0].set(key, y); nodes[1].set(key, x)
        self.mutate('ppt/slides/slide1.xml', swap)
        self.assertFalse(compare(self.before, self.after, 'layout', [1])['passed'])

    def test_layout_preserves_inherited_text_and_allows_geometry(self):
        for part in ('ppt/slideLayouts/slideLayout7.xml', 'ppt/slideMasters/slideMaster1.xml'):
            with self.subTest(part=part):
                self.after.write_bytes(self.before.read_bytes())
                def move(root):
                    root.xpath('//a:off', namespaces=N)[0].set('x', '999999')
                self.mutate(part, move)
                self.assertTrue(compare(self.before, self.after, 'layout')['passed'])
                def edit(root):
                    root.xpath('//a:t', namespaces=N)[0].text = 'Changed inherited claim'
                self.mutate(part, edit)
                self.assertFalse(compare(self.before, self.after, 'layout')['passed'])

    def test_layout_rejects_notes_backlink_change(self):
        def edit(root):
            for rel in root:
                if rel.get('Type', '').endswith('/slide'):
                    rel.set('Target', '../slides/slide2.xml')
        self.mutate('ppt/notesSlides/_rels/notesSlide1.xml.rels', edit)
        self.assertFalse(compare(self.before, self.after, 'layout', [1])['passed'])

    def test_layout_rejects_custom_notes_text_change(self):
        def custom(root):
            body = root.xpath('//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]', namespaces=N)[0]
            placeholder = body.xpath('.//p:ph', namespaces=N)[0]
            placeholder.getparent().remove(placeholder)
        self.mutate('ppt/notesSlides/notesSlide1.xml', custom, source=self.before, destination=self.before)
        self.after.write_bytes(self.before.read_bytes())
        def edit(root):
            root.xpath('//a:t', namespaces=N)[0].text = 'Changed custom note'
        self.mutate('ppt/notesSlides/notesSlide1.xml', edit)
        self.assertFalse(compare(self.before, self.after, 'layout', [1])['passed'])

    def test_layout_rejects_automatic_transition_and_timing(self):
        for name in ('transition', 'timing'):
            with self.subTest(name=name):
                self.after.write_bytes(self.before.read_bytes())
                def edit(root):
                    node = etree.SubElement(root, '{' + N['p'] + '}' + name)
                    if name == 'transition':
                        node.set('advTm', '1')
                    else:
                        etree.SubElement(node, '{' + N['p'] + '}tnLst')
                self.mutate('ppt/slides/slide1.xml', edit)
                self.assertFalse(compare(self.before, self.after, 'layout', [1])['passed'])

    def test_reorder_rejects_id_changes_pair_swaps_and_list_attributes(self):
        def id_change(root):
            root.find('p:sldIdLst/p:sldId', N).set('id', '9999')
        def pair_swap(root):
            nodes = list(root.find('p:sldIdLst', N))
            x, y = nodes[0].get('id'), nodes[1].get('id')
            nodes[0].set('id', y); nodes[1].set('id', x)
        def list_attr(root):
            root.find('p:sldIdLst', N).set('unexpected', '1')
        for edit in (id_change, pair_swap, list_attr):
            with self.subTest(edit=edit.__name__):
                self.after.write_bytes(self.before.read_bytes())
                self.mutate('ppt/presentation.xml', edit)
                self.assertFalse(compare(self.before, self.after, 'reorder')['passed'])

    def test_layout_allows_run_split_but_preserves_breaks(self):
        def split(root):
            from copy import deepcopy
            run = root.xpath('//a:r', namespaces=N)[0]
            clone = deepcopy(run)
            text = run.find('a:t', N).text
            run.find('a:t', N).text = text[:8]
            clone.find('a:t', N).text = text[8:]
            run.addnext(clone)
        self.mutate('ppt/slides/slide1.xml', split)
        self.assertTrue(compare(self.before, self.after, 'layout', [1])['passed'])
        def add_break(root):
            run = root.xpath('//a:r', namespaces=N)[0]
            run.addnext(etree.Element('{' + N['a'] + '}br'))
        self.mutate('ppt/slides/slide1.xml', add_break)
        self.assertFalse(compare(self.before, self.after, 'layout', [1])['passed'])

    def test_layout_preserves_hyperlink_actions_and_linked_spans(self):
        prs = Presentation(self.before)
        run = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
        run.hyperlink.address = 'https://example.org/research'
        prs.save(self.before)
        self.after.write_bytes(self.before.read_bytes())
        def split(root):
            from copy import deepcopy
            run = root.xpath('//a:r', namespaces=N)[0]
            clone = deepcopy(run)
            text = run.find('a:t', N).text
            run.find('a:t', N).text = text[:8]
            clone.find('a:t', N).text = text[8:]
            run.addnext(clone)
        self.mutate('ppt/slides/slide1.xml', split)
        self.assertTrue(compare(self.before, self.after, 'layout', [1])['passed'])
        def action(root):
            root.xpath('//a:hlinkClick', namespaces=N)[0].set('action', 'ppaction://hlinkshowjump?jump=endshow')
        self.mutate('ppt/slides/slide1.xml', action)
        self.assertFalse(compare(self.before, self.after, 'layout', [1])['passed'])
        self.after.write_bytes(self.before.read_bytes())
        self.mutate('ppt/slides/slide1.xml', split)
        def remove_link(root):
            node = root.xpath('//a:hlinkClick', namespaces=N)[0]
            node.getparent().remove(node)
        self.mutate('ppt/slides/slide1.xml', remove_link)
        self.assertFalse(compare(self.before, self.after, 'layout', [1])['passed'])

    def test_cli_returns_nonzero_and_machine_readable_report(self):
        self.mutate("ppt/slides/slide2.xml", lambda root: root.xpath('//a:srgbClr', namespaces=N)[0].set("val", "275A92"))
        run = subprocess.run([sys.executable, str(ROOT / "skills/pptx/scripts/compare_decks.py"), str(self.before), str(self.after), "--mode", "palette", "--slides", "1"], capture_output=True, text=True)
        import json
        self.assertEqual(run.returncode, 1)
        self.assertFalse(json.loads(run.stdout)["passed"])


if __name__ == "__main__":
    unittest.main()
