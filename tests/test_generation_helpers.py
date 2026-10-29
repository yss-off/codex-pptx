"""Geometry, image sizing, and notes regressions for actual generation helpers."""
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

from lxml import etree
from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / 'skills/pptx/scripts'
sys.path.insert(0, str(SCRIPTS / 'office'))
from validators.pptx import PPTXSchemaValidator

HELPERS = SCRIPTS / 'pptx_helpers.js'
NODE = shutil.which('node')
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
EMU = 914400


def node_env():
    env = os.environ.copy()
    candidates = [SCRIPTS.parent / 'node_modules']
    if shutil.which('npm'):
        result = subprocess.run(['npm', 'root', '-g'], capture_output=True, text=True)
        if result.returncode == 0:
            candidates.append(Path(result.stdout.strip()))
    existing = env.get('NODE_PATH')
    env['NODE_PATH'] = os.pathsep.join([str(p) for p in candidates] + ([existing] if existing else []))
    return env


@unittest.skipUnless(NODE, 'Node.js is required for generation-helper tests')
class HelperGeometryTests(unittest.TestCase):
    def run_js(self, body):
        result = subprocess.run([NODE, '-e', f'const h = require({json.dumps(str(HELPERS))});\n' + body],
                                check=True, capture_output=True, text=True)
        return json.loads(result.stdout)

    def test_contain_and_crop_preserve_source_aspect(self):
        values = self.run_js('console.log(JSON.stringify([h.imageContain(2400,1000,{x:1,y:2,w:4,h:4}),h.imageCrop(2400,1000,{x:1,y:2,w:4,h:4})]));')
        contain, crop = values
        self.assertAlmostEqual(contain['w'] / contain['h'], 2.4)
        self.assertAlmostEqual(contain['w'], 4)
        self.assertAlmostEqual(contain['y'] + contain['h'] / 2, 4)
        self.assertAlmostEqual(crop['w'] / crop['h'], 2.4)
        self.assertEqual(crop['sizing'], {'type': 'cover', 'w': 4, 'h': 4})

    def test_bad_image_dimensions_and_lines_rejected(self):
        failures = self.run_js('''
const bad = [
 () => h.imageContain(0,1,{x:0,y:0,w:2,h:2}),
 () => h.imageCrop(2,1,{x:0,y:0,w:-2,h:2}),
 () => h.imageContain(2,1,{x:0,y:0,w:Infinity,h:2}),
 () => h.imageCrop(NaN,1,{x:0,y:0,w:2,h:2}),
 () => h.lineBetween(1,1,1,1),
 () => h.lineBetween(0,0,1,1,{rotate:90}),
 () => h.lineBetween('0',0,1,1)
];
console.log(JSON.stringify(bad.map(fn => { try { fn(); return false; } catch (_) { return true; } })));
''')
        self.assertEqual(failures, [True] * 7)

    def test_all_line_directions_and_geometry_override(self):
        starts_ends = [[2,2,4,2], [4,2,2,2], [2,2,2,4], [2,4,2,2],
                       [2,2,4,4], [4,4,2,2], [2,4,4,2], [4,2,2,4]]
        result = self.run_js(f'''const pts={json.dumps(starts_ends)};
console.log(JSON.stringify(pts.map(p => h.lineBetween(...p, {{x:99,w:-10,flipH:true,line:{{endArrowType:'triangle'}}}}))));''')
        for expected, geom in zip(starts_ends, result):
            self.assertGreaterEqual(geom['w'], 0)
            self.assertGreaterEqual(geom['h'], 0)
            x, y, w, height = (geom[k] for k in ('x', 'y', 'w', 'h'))
            actual = [x + w if geom['flipH'] else x, y + height if geom['flipV'] else y,
                      x if geom['flipH'] else x + w, y if geom['flipV'] else y + height]
            self.assertEqual(actual, expected)
            self.assertEqual(geom['line']['endArrowType'], 'triangle')


@unittest.skipUnless(NODE, 'Node.js is required for PptxGenJS integration tests')
class GeneratedDeckTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.env = node_env()
        probe = subprocess.run([NODE, '-e', "require('pptxgenjs')"], env=cls.env, capture_output=True)
        if probe.returncode:
            raise unittest.SkipTest('PptxGenJS is not installed; pure geometry tests still run')

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.deck = self.root / 'generated.pptx'
        image = self.root / 'wide.png'
        Image.new('RGB', (240,100), 'blue').save(image)
        self.endpoints = [[1,1,3,1], [3,2,1,2], [4,1,4,3], [5,3,5,1],
                          [6,1,8,3], [8,4,6,2], [9,3,11,1], [11,2,9,4]]
        script = self.root / 'generate.js'
        script.write_text(f'''
const PptxGenJS=require('pptxgenjs');
const h=require({json.dumps(str(HELPERS))});
(async()=>{{
 const p=new PptxGenJS();p.layout='LAYOUT_WIDE';
 const a=p.addSlide();a.addText('geometry',{{x:0,y:0,w:3,h:0.5}});
 for(const pts of {json.dumps(self.endpoints)}) a.addShape(p.ShapeType.line,h.lineBetween(...pts,{{line:{{color:'123456',endArrowType:'triangle'}}}}));
 a.addNotes('Notes for geometry');
 const b=p.addSlide();b.addText('images',{{x:0,y:0,w:3,h:0.5}});
 b.addImage({{path:{json.dumps(str(image))},...h.imageContain(240,100,{{x:1,y:1,w:4,h:4}})}});
 b.addImage({{path:{json.dumps(str(image))},...h.imageCrop(240,100,{{x:6,y:1,w:4,h:4}})}});
 b.addNotes('Notes for images');
 await p.writeFile({{fileName:{json.dumps(str(self.deck))}}});
}})().catch(e=>{{console.error(e);process.exit(1)}});
''')
        subprocess.run([NODE, str(script)], env=self.env, check=True, capture_output=True)
        self.unpacked = self.root / 'unpacked'
        with zipfile.ZipFile(self.deck) as archive:
            archive.extractall(self.unpacked)

    def test_generated_lines_have_positive_extents_and_correct_arrow_ends(self):
        root = etree.parse(str(self.unpacked / 'ppt/slides/slide1.xml'))
        lines = root.xpath('//p:sp[p:spPr/a:prstGeom[@prst="line"]]', namespaces=NS)
        self.assertEqual(len(lines), len(self.endpoints))
        for shape, pts in zip(lines, self.endpoints):
            transform = shape.find('p:spPr/a:xfrm', NS)
            off, ext = transform.find('a:off', NS), transform.find('a:ext', NS)
            x,y,w,h = [int(v)/EMU for v in (off.get('x'),off.get('y'),ext.get('cx'),ext.get('cy'))]
            self.assertGreaterEqual(w, 0); self.assertGreaterEqual(h, 0)
            hx = transform.get('flipH') in ('1','true')
            vy = transform.get('flipV') in ('1','true')
            actual = [x+w if hx else x, y+h if vy else y, x if hx else x+w, y if vy else y+h]
            self.assertEqual(actual, pts)
            self.assertEqual(shape.find('p:spPr/a:ln/a:tailEnd', NS).get('type'), 'triangle')
        validator = PPTXSchemaValidator(self.unpacked)
        valid, errors = validator._validate_single_file_xsd(self.unpacked / 'ppt/slides/slide1.xml', self.unpacked)
        self.assertTrue(valid, errors)

    def test_generated_images_preserve_ratio_or_crop_without_stretch(self):
        root = etree.parse(str(self.unpacked / 'ppt/slides/slide2.xml'))
        pics = root.xpath('//p:pic', namespaces=NS)
        ext = pics[0].find('p:spPr/a:xfrm/a:ext', NS)
        self.assertAlmostEqual(int(ext.get('cx'))/int(ext.get('cy')), 2.4, places=5)
        self.assertIsNone(pics[0].find('p:blipFill/a:srcRect', NS))
        ext = pics[1].find('p:spPr/a:xfrm/a:ext', NS)
        self.assertEqual((int(ext.get('cx')), int(ext.get('cy'))), (4*EMU,4*EMU))
        crop = pics[1].find('p:blipFill/a:srcRect', NS)
        self.assertAlmostEqual(int(crop.get('l'))/100000, (1-1/2.4)/2, places=4)
        self.assertEqual(crop.get('l'),crop.get('r'))
        self.assertEqual((crop.get('t'),crop.get('b')), ('0','0'))
        validator = PPTXSchemaValidator(self.unpacked)
        valid, errors = validator._validate_single_file_xsd(self.unpacked / 'ppt/slides/slide2.xml', self.unpacked)
        self.assertTrue(valid, errors)

    def test_crop_exports_long_images_across_unit_threshold(self):
        # A 30:1 image covered into 4x4 used to supply virtual width=120;
        # PptxGenJS interpreted that as EMU while interpreting height=4 as inches.
        cases = [
            (3000, 100, (48333, 48333, 0, 0)),
            (100, 3000, (0, 0, 48333, 48333)),
            (2500, 100, (48000, 48000, 0, 0)),  # old virtual width exactly 100
            (100, 2500, (0, 0, 48000, 48000)),
            (2400, 1000, (29167, 29167, 0, 0)),
            (1000, 2400, (0, 0, 29167, 29167)),
        ]
        inputs = []
        for i, (width, height, _) in enumerate(cases):
            image = self.root / f'crop-{i}.png'
            Image.new('RGB', (width, height), 'blue').save(image)
            inputs.append([str(image), width, height])
        deck = self.root / 'crop-thresholds.pptx'
        script = f'''
const PptxGenJS=require('pptxgenjs');
const h=require({json.dumps(str(HELPERS))});
(async()=>{{
 const p=new PptxGenJS();p.layout='LAYOUT_WIDE';
 for (const [path,w,hgt] of {json.dumps(inputs)}) {{
  p.addSlide().addImage({{path,...h.imageCrop(w,hgt,{{x:1,y:1,w:4,h:4}})}});
 }}
 await p.writeFile({{fileName:{json.dumps(str(deck))}}});
}})().catch(e=>{{console.error(e);process.exit(1)}});
'''
        subprocess.run([NODE, '-e', script], env=self.env, check=True, capture_output=True)
        with zipfile.ZipFile(deck) as archive:
            for i, (width, height, expected) in enumerate(cases, 1):
                with self.subTest(source=(width, height)):
                    root = etree.fromstring(archive.read(f'ppt/slides/slide{i}.xml'))
                    pic = root.find('.//p:pic', NS)
                    crop = pic.find('p:blipFill/a:srcRect', NS)
                    for side, value in zip(('l', 'r', 't', 'b'), expected):
                        self.assertAlmostEqual(int(crop.get(side)), value, delta=1)
                    extent = pic.find('p:spPr/a:xfrm/a:ext', NS)
                    self.assertEqual((int(extent.get('cx')), int(extent.get('cy'))),
                                     (4*EMU, 4*EMU))
                    offset = pic.find('p:spPr/a:xfrm/a:off', NS)
                    self.assertEqual((int(offset.get('x')), int(offset.get('y'))), (EMU, EMU))

    def test_notes_follow_slides_after_reordering(self):
        prs = Presentation(self.deck)
        self.assertEqual([s.notes_slide.notes_text_frame.text for s in prs.slides], ['Notes for geometry','Notes for images'])
        # Change display order only; physical slide parts keep their own notes.
        path = self.unpacked / 'ppt/presentation.xml'
        tree = etree.parse(str(path)); ids = tree.getroot().find('p:sldIdLst', NS)
        ids.insert(0, ids[-1]); tree.write(str(path), encoding='UTF-8', xml_declaration=True)
        reordered = self.root / 'reordered.pptx'
        with zipfile.ZipFile(reordered, 'w') as archive:
            for item in self.unpacked.rglob('*'):
                if item.is_file(): archive.write(item, item.relative_to(self.unpacked))
        prs = Presentation(reordered)
        self.assertEqual([s.notes_slide.notes_text_frame.text for s in prs.slides], ['Notes for images','Notes for geometry'])
        self.assertEqual([s.shapes[0].text for s in prs.slides], ['images','geometry'])

    def test_generated_notes_schema_issue_is_reported_not_assumed_valid(self):
        # PptxGenJS 4.0.1 emits the notes-master list after the slide list. Do not
        # silently suppress this exporter problem or rewrite its internals.
        path = self.unpacked / 'ppt/presentation.xml'
        validator = PPTXSchemaValidator(self.unpacked)
        valid, errors = validator._validate_single_file_xsd(path, self.unpacked)
        names = [etree.QName(n).localname for n in etree.parse(str(path)).getroot()]
        if names.index('notesMasterIdLst') > names.index('sldIdLst'):
            self.assertFalse(valid)
            self.assertTrue(any('notesMasterIdLst' in e for e in errors), errors)
        else:
            self.assertTrue(valid, errors)


if __name__ == '__main__':
    unittest.main()
