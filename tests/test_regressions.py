"""Behavioral regressions for the distributable PPTX workflows."""

import contextlib
import io
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills/pptx/scripts"
sys.path[:0] = [str(SCRIPTS), str(SCRIPTS / "office")]
import add_slide
import clean
import thumbnail
from office import soffice
from office import pack
from office.validators.pptx import PPTXSchemaValidator
from office.opc import resolve_target

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/package/2006/relationships"


class WorkflowRegressions(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.original = self.root / "source.pptx"
        prs = Presentation()
        for index in range(3):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = f"Slide {index}"
        prs.slides[1]._element.set("show", "0")
        data = CategoryChartData()
        data.categories = ["Q1", "Q2"]
        data.add_series("Sales", [10, 20])
        chart = prs.slides[0].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(5), Inches(3), data)
        # python-pptx emits signed axis IDs; this positive fixture must obey the
        # unsigned OOXML schema rather than depend on an exporter defect.
        for node in chart.chart._chartSpace.iter():
            if node.tag.rsplit("}", 1)[-1] in {"axId", "crossAx"}:
                node.set("val", str(int(node.get("val")) & 0xFFFFFFFF))
        prs.save(self.original)
        self.unpacked = self.root / "unpacked"
        with zipfile.ZipFile(self.original) as archive:
            archive.extractall(self.unpacked)

    def rewrite(self, relative, mutate):
        path = self.unpacked / relative
        tree = etree.parse(str(path))
        mutate(tree.getroot())
        tree.write(str(path), encoding="UTF-8", xml_declaration=True)

    def test_clean_preserves_absolute_chart_and_workbook_targets(self):
        rel = "ppt/slides/_rels/slide1.xml.rels"
        self.rewrite(rel, lambda root: [r.set("Target", "/ppt/charts/chart1.xml") for r in root if r.get("Type", "").endswith("/chart")])
        clean.clean_unused_files(self.unpacked)
        self.assertTrue((self.unpacked / "ppt/charts/chart1.xml").exists())
        self.assertTrue(list((self.unpacked / "ppt/embeddings").glob("*.xlsx")))

    def test_clean_preserves_namespace_independent_slide_list(self):
        path = self.unpacked / "ppt/presentation.xml"
        path.write_text(path.read_text().replace("xmlns:p=", "xmlns:pres=").replace("<p:", "<pres:").replace("</p:", "</pres:"))
        clean.clean_unused_files(self.unpacked)
        self.assertEqual(len(list((self.unpacked / "ppt/slides").glob("slide*.xml"))), 3)

    def test_clean_preserves_absolute_slide_targets(self):
        self.rewrite("ppt/_rels/presentation.xml.rels", lambda root: [r.set("Target", "/ppt/" + r.get("Target")) for r in root if r.get("Type", "").endswith("/slide")])
        clean.clean_unused_files(self.unpacked)
        self.assertEqual(len(list((self.unpacked / "ppt/slides").glob("slide*.xml"))), 3)

    def test_clean_still_removes_deleted_slide(self):
        self.rewrite("ppt/presentation.xml", lambda root: root.find(f"{{{P}}}sldIdLst").remove(root.find(f"{{{P}}}sldIdLst")[-1]))
        clean.clean_unused_files(self.unpacked)
        self.assertFalse((self.unpacked / "ppt/slides/slide3.xml").exists())
        self.assertTrue((self.unpacked / "ppt/slides/slide1.xml").exists())

    def test_invalid_slide_layout_and_master_fail_xsd(self):
        validator = PPTXSchemaValidator(self.unpacked)
        for relative in ["ppt/slides/slide1.xml", "ppt/slideLayouts/slideLayout1.xml", "ppt/slideMasters/slideMaster1.xml"]:
            with self.subTest(part=relative):
                self.rewrite(relative, lambda root: root.insert(0, etree.Element(f"{{{P}}}invalidElement")))
                valid, errors = validator._validate_single_file_xsd(self.unpacked / relative, self.unpacked)
                self.assertIs(valid, False)
                self.assertTrue(any("invalidElement" in e for e in errors))

    def test_pack_rejects_new_invalid_slide(self):
        self.rewrite("ppt/slides/slide1.xml", lambda root: root.insert(0, etree.Element(f"{{{P}}}invalidElement")))
        output = self.root / "invalid.pptx"
        with contextlib.redirect_stdout(io.StringIO()):
            _, message = pack.pack(str(self.unpacked), str(output), str(self.original))
        self.assertIn("Error", message)
        self.assertFalse(output.exists())

    def test_pack_without_original_rejects_dangling_reference(self):
        self.rewrite("ppt/slides/_rels/slide1.xml.rels", lambda root: etree.SubElement(root, f"{{{R}}}Relationship", Id="rId999", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image", Target="../media/missing.png"))
        output = self.root / "invalid.pptx"
        with contextlib.redirect_stdout(io.StringIO()):
            _, message = pack.pack(str(self.unpacked), str(output))
        self.assertIn("Error", message)
        self.assertFalse(output.exists())

    def test_pack_rejects_missing_original(self):
        _, message = pack.pack(str(self.unpacked), str(self.root / "out.pptx"), str(self.root / "missing.pptx"))
        self.assertIn("Error", message)

    def test_pack_valid_deck_without_original(self):
        output = self.root / "valid.pptx"
        with contextlib.redirect_stdout(io.StringIO()):
            _, message = pack.pack(str(self.unpacked), str(output))
        self.assertNotIn("Error", message)
        self.assertEqual(len(Presentation(output).slides), 3)

    def test_hidden_slide_gets_placeholder_and_keeps_order(self):
        info = thumbnail.get_slide_info(self.original)
        self.assertEqual([s["hidden"] for s in info], [False, True, False])
        images = []
        for i in range(2):
            image = self.root / f"visible-{i}.jpg"
            Image.new("RGB", (20, 10)).save(image)
            images.append(image)
        slides = thumbnail.build_slide_list(info, images, self.root)
        self.assertEqual([s[1] for s in slides], ["slide1.xml", "slide2.xml (hidden)", "slide3.xml"])
        self.assertEqual(slides[-1][0], images[-1])

    def test_duplicate_chart_and_workbook_are_independent(self):
        with contextlib.redirect_stdout(io.StringIO()):
            add_slide.duplicate_slide(self.unpacked, "slide1.xml")
        def chart_target(slide):
            tree = etree.parse(str(self.unpacked / f"ppt/slides/_rels/{slide}.xml.rels"))
            return next(r.get("Target") for r in tree.getroot() if r.get("Type", "").endswith("/chart"))
        self.assertNotEqual(chart_target("slide1"), chart_target("slide4"))
        charts = list((self.unpacked / "ppt/charts").glob("chart*.xml"))
        self.assertEqual(len(charts), 2)
        workbooks = list((self.unpacked / "ppt/embeddings").glob("*.xlsx"))
        self.assertEqual(len(workbooks), 2)
        self.assertEqual(workbooks[0].read_bytes(), workbooks[1].read_bytes())
        # Materially changing the copied chart must leave the source untouched.
        original_bytes = (self.unpacked / "ppt/charts/chart1.xml").read_bytes()
        copy = next(p for p in charts if p.name != "chart1.xml")
        copy.write_bytes(copy.read_bytes().replace(b">10<", b">999<"))
        self.assertEqual((self.unpacked / "ppt/charts/chart1.xml").read_bytes(), original_bytes)

    def test_duplicate_chart_survives_clean_validate_pack(self):
        with contextlib.redirect_stdout(io.StringIO()):
            add_slide.duplicate_slide(self.unpacked, "slide1.xml")
        rels = etree.parse(str(self.unpacked / "ppt/_rels/presentation.xml.rels"))
        rid = next(r.get("Id") for r in rels.getroot() if r.get("Target") == "slides/slide4.xml")
        def append_slide(root):
            etree.SubElement(root.find(f"{{{P}}}sldIdLst"), f"{{{P}}}sldId", id="259", attrib={"{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id": rid})
        self.rewrite("ppt/presentation.xml", append_slide)
        chart = self.unpacked / "ppt/charts/chart2.xml"
        chart.write_bytes(chart.read_bytes().replace(b">10<", b">999<"))
        clean.clean_unused_files(self.unpacked)
        output = self.root / "duplicated.pptx"
        with contextlib.redirect_stdout(io.StringIO()):
            _, message = pack.pack(str(self.unpacked), str(output), str(self.original))
        self.assertNotIn("Error", message)
        prs = Presentation(output)
        values = lambda slide: next(shape.chart for shape in slide.shapes if shape.has_chart).series[0].values
        self.assertEqual(tuple(values(prs.slides[0])), (10, 20))
        self.assertEqual(tuple(values(prs.slides[3])), (999, 20))

    def test_opc_targets_are_contained_and_external_links_ignored(self):
        rels = self.unpacked / "ppt/slides/_rels/slide1.xml.rels"
        self.assertIsNone(resolve_target(self.unpacked, rels, "../../../../outside.xml"))
        self.assertIsNone(resolve_target(self.unpacked, rels, "file:///tmp/outside.xml", "External"))
        self.assertIsNone(resolve_target(self.unpacked, rels, "https://example.com/chart.xml"))
        self.assertEqual(resolve_target(self.unpacked, rels, "../media/my%20image.png"), self.unpacked / "ppt/media/my image.png")

    def test_pack_skip_requires_explicit_opt_out(self):
        self.rewrite("ppt/slides/slide1.xml", lambda root: root.insert(0, etree.Element(f"{{{P}}}invalidElement")))
        output = self.root / "explicitly-unchecked.pptx"
        _, message = pack.pack(str(self.unpacked), str(output), validate=False)
        self.assertNotIn("Error", message)
        self.assertTrue(output.exists())

    def test_editing_instructions_use_available_codex_tools(self):
        guide = (ROOT / "skills/pptx/editing.md").read_text()
        self.assertNotIn("Use the Edit tool", guide)
        self.assertIn("apply_patch", guide)

    def test_shim_ignores_preexisting_shared_library(self):
        untrusted = self.root / "lo_socket_shim.so"
        untrusted.write_text("untrusted")
        def fake_compile(args, **kwargs):
            Path(args[args.index("-o") + 1]).write_bytes(b"newly compiled")
        with patch.object(soffice.tempfile, "gettempdir", return_value=str(self.root)), patch.object(soffice.subprocess, "run", side_effect=fake_compile):
            shim = soffice._ensure_shim()
        self.assertNotEqual(shim, untrusted)
        self.assertEqual(shim.read_bytes(), b"newly compiled")
        self.assertEqual(shim.parent.stat().st_mode & 0o777, 0o700)


if __name__ == "__main__":
    unittest.main()
